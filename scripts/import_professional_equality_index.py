"""Index the official professional equality declarations by exact SIREN.

Only aggregate company/UES results are retained. Employee records, pay amounts,
contacts and source legal names are never copied into the local index.
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path
import re
import sqlite3

from openpyxl import load_workbook
import requests

try:
    from scripts.import_public_grants import compact_text, load_targets, normalize_header
except ModuleNotFoundError:  # Direct execution with scripts/ as sys.path[0].
    from import_public_grants import compact_text, load_targets, normalize_header  # type: ignore[no-redef]


ROOT = Path(__file__).resolve().parents[1]
DATASET_ID = "index-egalite-professionnelle-f-h-des-entreprises-de-50-salaries-ou-plus"
DATASET_API = f"https://www.data.gouv.fr/api/1/datasets/{DATASET_ID}/"
DATASET_URL = f"https://www.data.gouv.fr/datasets/{DATASET_ID}"
EGAPRO_URL = "https://egapro.travail.gouv.fr/consulter-index"
LICENSE_NAME = "Licence Ouverte 2.0"
LICENSE_URL = "https://www.etalab.gouv.fr/licence-ouverte-open-licence/"
DEFAULT_COMPANY_DB = ROOT / "data/guadeloupe-enterprises.sqlite"
DEFAULT_CACHE_DIR = ROOT / "data/imports/professional-equality"
DEFAULT_OUTPUT = ROOT / "data/professional-equality-index.sqlite"
USER_AGENT = "guadeloupe-entreprises-professional-equality/0.1 (+public open-data reuse)"
MAX_DOWNLOAD_BYTES = 50 * 1024 * 1024
RESOURCE_TITLE = "Index Egalité Professionnelle F/H"

SCORE_FIELDS = {
    "pay_gap": ("noteecartremuneration", 40),
    "raise_gap_no_promotion": ("noteecarttauxdaugmentationhorspromotion", 20),
    "promotion_gap": ("noteecarttauxdepromotion", 15),
    "raise_gap": ("noteecarttauxdaugmentation", 35),
    "maternity_return": ("noteretourcongematernite", 15),
    "highest_remuneration": ("notehautesremunerations", 10),
    "index": ("noteindex", 100),
}


def source_metadata(session: requests.Session) -> dict[str, object]:
    response = session.get(DATASET_API, timeout=30)
    response.raise_for_status()
    payload = response.json()
    resource = next(
        (
            item for item in payload.get("resources", [])
            if isinstance(item, dict)
            and str(item.get("title") or "") == RESOURCE_TITLE
            and str(item.get("format") or "").lower() == "xlsx"
        ),
        None,
    )
    if resource is None:
        raise RuntimeError("Ressource XLSX de l'Index Egapro introuvable")
    if str(payload.get("license") or "").lower() not in {"fr-lo", "lov2"}:
        raise RuntimeError("Licence ouverte de l'Index Egapro non confirmée")
    return {
        "dataset_id": str(payload.get("id") or DATASET_ID),
        "title": compact_text(payload.get("title"), 500),
        "source_updated_at": str(payload.get("last_update") or payload.get("last_modified") or ""),
        "license_code": str(payload.get("license") or ""),
        "license_name": LICENSE_NAME,
        "license_url": LICENSE_URL,
        "resource": {
            "id": str(resource.get("id") or ""),
            "title": str(resource.get("title") or RESOURCE_TITLE),
            "url": str(resource.get("url") or ""),
            "last_modified": str(resource.get("last_modified") or ""),
        },
    }


def download_resource(
    session: requests.Session,
    resource: dict[str, object],
    cache_dir: Path,
    refresh: bool,
) -> tuple[Path, bool]:
    target = cache_dir / f"{resource['id']}.xlsx"
    sidecar = cache_dir / f"{resource['id']}.metadata.json"
    signature = {"url": resource.get("url"), "last_modified": resource.get("last_modified")}
    if not refresh and target.exists() and sidecar.exists():
        try:
            if json.loads(sidecar.read_text(encoding="utf-8")) == signature and target.stat().st_size > 0:
                return target, False
        except (OSError, json.JSONDecodeError):
            pass
    cache_dir.mkdir(parents=True, exist_ok=True)
    partial = target.with_suffix(".xlsx.part")
    partial.unlink(missing_ok=True)
    downloaded = 0
    prefix = b""
    with session.get(str(resource["url"]), stream=True, timeout=(30, 180)) as response:
        response.raise_for_status()
        with partial.open("wb") as output:
            for chunk in response.iter_content(256 * 1024):
                if not chunk:
                    continue
                if not prefix:
                    prefix = chunk[:4]
                    if not prefix.startswith(b"PK"):
                        raise ValueError("La ressource Egapro n'est pas un classeur XLSX")
                downloaded += len(chunk)
                if downloaded > MAX_DOWNLOAD_BYTES:
                    partial.unlink(missing_ok=True)
                    raise ValueError("Ressource Egapro trop volumineuse")
                output.write(chunk)
    partial.replace(target)
    sidecar.write_text(json.dumps(signature, ensure_ascii=False, indent=2), encoding="utf-8")
    return target, True


def normalized_siren(value: object) -> str | None:
    digits = re.sub(r"\D", "", compact_text(value, 64) or "")
    return digits if len(digits) == 9 else None


def constituent_sirens(value: object) -> set[str]:
    return set(re.findall(r"(?<!\d)(\d{9})(?!\d)", compact_text(value, 20_000) or ""))


def parse_year(value: object) -> int | None:
    try:
        year = int(value)
    except (TypeError, ValueError):
        return None
    return year if 2018 <= year <= datetime.now().year else None


def parse_score(value: object, maximum: int) -> tuple[int | None, str]:
    if value is None or compact_text(value, 20) is None:
        return None, "not_applicable"
    raw = str(value).strip()
    if normalize_header(raw) in {"nc", "noncalculable"}:
        return None, "not_calculable"
    try:
        score = int(float(raw.replace(",", ".")))
    except ValueError:
        return None, "invalid"
    if score < 0 or score > maximum:
        return None, "invalid"
    return score, "calculated"


def split_naf(value: object) -> tuple[str | None, str | None]:
    raw = compact_text(value, 500)
    if not raw:
        return None, None
    match = re.match(r"^([0-9]{2}\.[0-9]{2}[A-Z])\s*-\s*(.+)$", raw)
    return (match.group(1), compact_text(match.group(2), 400)) if match else (None, raw)


def declaration_fingerprint(profile: dict[str, object]) -> str:
    stable = (
        profile.get("siren"), profile.get("declaring_siren"), profile.get("reference_year"),
        profile.get("match_scope"), profile.get("structure_type"),
    )
    return sha256("|".join(str(value or "").casefold() for value in stable).encode("utf-8")).hexdigest()


def parse_resource(
    path: Path,
    resource: dict[str, object],
    company_db: Path,
) -> tuple[list[dict[str, object]], dict[str, object]]:
    targets = load_targets(company_db)
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    raw_headers = next(rows, None)
    if not raw_headers:
        raise ValueError("Classeur Egapro vide")
    positions = {normalize_header(value): index for index, value in enumerate(raw_headers)}
    required = {
        "annee", "structure", "tranchedeffectifs", "siren", "nomues", "entreprisesuessiren",
        "region", "departement", "pays", "codenaf", *(column for column, _ in SCORE_FIELDS.values()),
    }
    if not required.issubset(positions):
        raise ValueError(f"Colonnes Egapro manquantes: {sorted(required - positions.keys())}")

    stats: dict[str, object] = {
        "resource_id": resource["id"], "resource_title": resource["title"], "sheet": worksheet.title,
        "rows_read": 0, "rows_with_valid_identifier": 0, "invalid_identifier_rows": 0,
        "invalid_year_rows": 0, "matched_source_rows": 0, "matched_links": 0,
        "direct_declarant_links": 0, "ues_member_links": 0, "rows_with_invalid_score": 0,
        "guadeloupe_declaration_links": 0, "outside_guadeloupe_declaration_links": 0,
        "calculable_index_links": 0, "non_calculable_index_links": 0,
    }
    profiles: list[dict[str, object]] = []
    for row_number, row in enumerate(rows, start=2):
        stats["rows_read"] = int(stats["rows_read"]) + 1
        declaring_siren = normalized_siren(row[positions["siren"]])
        if not declaring_siren:
            stats["invalid_identifier_rows"] = int(stats["invalid_identifier_rows"]) + 1
            continue
        stats["rows_with_valid_identifier"] = int(stats["rows_with_valid_identifier"]) + 1
        year = parse_year(row[positions["annee"]])
        if year is None:
            stats["invalid_year_rows"] = int(stats["invalid_year_rows"]) + 1
            continue

        matches: dict[str, str] = {}
        if declaring_siren in targets.company_sirens:
            matches[declaring_siren] = "exact_declarant"
        for member_siren in constituent_sirens(row[positions["entreprisesuessiren"]]):
            if member_siren in targets.company_sirens and member_siren not in matches:
                matches[member_siren] = "ues_member"
        if not matches:
            continue
        stats["matched_source_rows"] = int(stats["matched_source_rows"]) + 1

        scores: dict[str, object] = {}
        invalid_score = False
        for field, (header, maximum) in SCORE_FIELDS.items():
            score, status = parse_score(row[positions[header]], maximum)
            scores[f"{field}_score"] = score
            scores[f"{field}_status"] = status
            invalid_score = invalid_score or status == "invalid"
        if invalid_score:
            stats["rows_with_invalid_score"] = int(stats["rows_with_invalid_score"]) + 1

        raw_structure = compact_text(row[positions["structure"]], 100)
        structure_type = "ues" if "economique" in normalize_header(raw_structure) else "company"
        region = compact_text(row[positions["region"]], 150)
        department = compact_text(row[positions["departement"]], 150)
        location_scope = "guadeloupe" if normalize_header(department) == "guadeloupe" else "outside_guadeloupe"
        naf_code, naf_label = split_naf(row[positions["codenaf"]])
        all_members = constituent_sirens(row[positions["entreprisesuessiren"]])

        for target_siren, match_scope in matches.items():
            profile: dict[str, object] = {
                "source_row_number": row_number,
                "siren": target_siren,
                "declaring_siren": declaring_siren,
                "match_scope": match_scope,
                "match_confidence": 1.0 if match_scope == "exact_declarant" else 0.99,
                "reference_year": year,
                "structure_type": structure_type,
                "workforce_band": compact_text(row[positions["tranchedeffectifs"]], 100),
                "ues_name": compact_text(row[positions["nomues"]], 300) if structure_type == "ues" else None,
                "ues_member_count": len(all_members) if structure_type == "ues" else None,
                "declaration_location_scope": location_scope,
                "declaring_region": region,
                "declaring_department": department,
                "declaring_country": compact_text(row[positions["pays"]], 100),
                "naf_code": naf_code,
                "naf_label": naf_label,
                **scores,
                "resource_id": resource["id"],
                "resource_title": resource["title"],
                "resource_url": resource["url"],
                "resource_last_modified": resource.get("last_modified") or None,
            }
            profile["fingerprint"] = declaration_fingerprint(profile)
            profiles.append(profile)
            stats["matched_links"] = int(stats["matched_links"]) + 1
            scope_counter = "direct_declarant_links" if match_scope == "exact_declarant" else "ues_member_links"
            stats[scope_counter] = int(stats[scope_counter]) + 1
            stats[f"{location_scope}_declaration_links"] = int(stats[f"{location_scope}_declaration_links"]) + 1
            index_status = str(scores["index_status"])
            if index_status == "calculated":
                stats["calculable_index_links"] = int(stats["calculable_index_links"]) + 1
            elif index_status == "not_calculable":
                stats["non_calculable_index_links"] = int(stats["non_calculable_index_links"]) + 1
    workbook.close()
    return profiles, stats


def build_index(
    profiles: list[dict[str, object]],
    output: Path,
    metadata: dict[str, object],
    stats: dict[str, object],
    downloaded: bool,
) -> dict[str, object]:
    partial = output.with_suffix(output.suffix + ".part")
    partial.unlink(missing_ok=True)
    output.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(partial)
    score_columns = ",\n".join(
        f"{field}_score INTEGER, {field}_status TEXT NOT NULL" for field in SCORE_FIELDS
    )
    connection.executescript(f"""
      CREATE TABLE professional_equality_declarations (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        fingerprint TEXT NOT NULL UNIQUE,
        source_row_number INTEGER NOT NULL,
        siren TEXT NOT NULL,
        declaring_siren TEXT NOT NULL,
        match_scope TEXT NOT NULL,
        match_confidence REAL NOT NULL,
        reference_year INTEGER NOT NULL,
        structure_type TEXT NOT NULL,
        workforce_band TEXT,
        ues_name TEXT,
        ues_member_count INTEGER,
        declaration_location_scope TEXT NOT NULL,
        declaring_region TEXT,
        declaring_department TEXT,
        declaring_country TEXT,
        naf_code TEXT,
        naf_label TEXT,
        {score_columns},
        resource_id TEXT NOT NULL,
        resource_title TEXT NOT NULL,
        resource_url TEXT NOT NULL,
        resource_last_modified TEXT,
        source_updated_at TEXT NOT NULL,
        dataset_url TEXT NOT NULL,
        egapro_url TEXT NOT NULL,
        license_name TEXT NOT NULL,
        license_url TEXT NOT NULL,
        imported_at TEXT NOT NULL
      );
      CREATE TABLE metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL);
      CREATE INDEX equality_siren_year_idx ON professional_equality_declarations(siren, reference_year DESC);
      CREATE INDEX equality_declarant_year_idx ON professional_equality_declarations(declaring_siren, reference_year DESC);
      CREATE INDEX equality_scope_idx ON professional_equality_declarations(match_scope, structure_type);
      CREATE INDEX equality_score_idx ON professional_equality_declarations(reference_year DESC, index_score);
    """)
    imported_at = datetime.now(timezone.utc).isoformat()
    source_updated_at = str(metadata.get("source_updated_at") or "")
    base_columns = (
        "fingerprint", "source_row_number", "siren", "declaring_siren", "match_scope", "match_confidence",
        "reference_year", "structure_type", "workforce_band", "ues_name", "ues_member_count",
        "declaration_location_scope", "declaring_region", "declaring_department", "declaring_country",
        "naf_code", "naf_label",
    )
    score_column_names = tuple(
        column for field in SCORE_FIELDS for column in (f"{field}_score", f"{field}_status")
    )
    resource_columns = ("resource_id", "resource_title", "resource_url", "resource_last_modified")
    columns = base_columns + score_column_names + resource_columns
    extra_values = (source_updated_at, DATASET_URL, EGAPRO_URL, metadata["license_name"], metadata["license_url"], imported_at)
    placeholders = ",".join("?" for _ in range(len(columns) + len(extra_values)))
    for profile in profiles:
        values = tuple(profile.get(column) for column in columns)
        connection.execute(
            f"INSERT OR IGNORE INTO professional_equality_declarations ({','.join(columns)}, source_updated_at, dataset_url, egapro_url, license_name, license_url, imported_at) VALUES ({placeholders})",
            values + extra_values,
        )
    aggregate = connection.execute(
        """SELECT COUNT(*), COUNT(DISTINCT siren),
                  SUM(match_scope = 'exact_declarant'), SUM(match_scope = 'ues_member'),
                  SUM(index_status = 'calculated'), SUM(index_status = 'not_calculable'),
                  SUM(declaration_location_scope = 'guadeloupe'),
                  MIN(reference_year), MAX(reference_year),
                  MIN(index_score), MAX(index_score)
           FROM professional_equality_declarations"""
    ).fetchone()
    report = {
        "source": metadata.get("title"),
        "source_url": DATASET_URL,
        "egapro_url": EGAPRO_URL,
        "source_updated_at": source_updated_at,
        "license": metadata.get("license_name"),
        "license_url": metadata.get("license_url"),
        "imported_at": imported_at,
        "downloaded_resource_count": int(downloaded),
        "cached_resource_count": int(not downloaded),
        "source_rows": int(stats["rows_read"]),
        "rows_with_valid_identifier": int(stats["rows_with_valid_identifier"]),
        "invalid_identifier_rows": int(stats["invalid_identifier_rows"]),
        "matched_source_rows": int(stats["matched_source_rows"]),
        "matched_declarations": aggregate[0],
        "matched_companies": aggregate[1],
        "direct_declarant_count": aggregate[2] or 0,
        "ues_member_count": aggregate[3] or 0,
        "calculable_index_count": aggregate[4] or 0,
        "non_calculable_index_count": aggregate[5] or 0,
        "guadeloupe_declaration_count": aggregate[6] or 0,
        "earliest_reference_year": aggregate[7],
        "latest_reference_year": aggregate[8],
        "minimum_calculable_index": aggregate[9],
        "maximum_calculable_index": aggregate[10],
        "resource_stats": stats,
        "contains_employee_records": False,
        "contains_pay_amounts": False,
        "contains_personal_names": False,
        "contains_contacts": False,
        "stores_source_legal_names": False,
        "join_uses_names": False,
        "database": str(output),
    }
    for key, value in report.items():
        if key in {"database", "resource_stats"}:
            continue
        encoded = json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list, bool)) else str(value or "")
        connection.execute("INSERT INTO metadata(key, value) VALUES (?, ?)", (key, encoded))
    connection.commit()
    connection.execute("ANALYZE")
    connection.close()
    partial.replace(output)
    output.with_suffix(".report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--company-db", type=Path, default=DEFAULT_COMPANY_DB)
    parser.add_argument("--cache-dir", type=Path, default=DEFAULT_CACHE_DIR)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--refresh", action="store_true")
    parser.add_argument("--skip-download", action="store_true")
    args = parser.parse_args()

    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT, "Accept": "application/json,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;q=0.9"})
    metadata = source_metadata(session)
    resource = metadata["resource"]
    if args.skip_download:
        path = args.cache_dir / f"{resource['id']}.xlsx"
        if not path.exists():
            raise SystemExit(f"Ressource Egapro absente du cache: {path}")
        downloaded = False
    else:
        path, downloaded = download_resource(session, resource, args.cache_dir, args.refresh)
    profiles, stats = parse_resource(path, resource, args.company_db)
    print(json.dumps(build_index(profiles, args.output, metadata, stats, downloaded), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
